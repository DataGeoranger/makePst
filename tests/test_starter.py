"""`makepst init` writes a workbook that builds a valid control file as it is."""
import os
import shlex
import sys

import openpyxl
import pandas as pd
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

from makepst import Pst, read_pst  # noqa: E402
from makepst.cli import main  # noqa: E402
from makepst.sections import ALL_SECTIONS, COMPUTED  # noqa: E402
from makepst.starter import DESCRIPTIONS  # noqa: E402


@pytest.fixture
def starter(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    main(['init', 'project.xlsx'])
    out = capsys.readouterr().out
    cmd = out.splitlines()[-1].strip()
    return tmp_path / 'project.xlsx', cmd


def test_init_builds_valid_pst(starter):
    book, cmd = starter
    assert cmd.startswith('makepst build project.pst estimation')
    main(shlex.split(cmd)[1:])
    p = read_pst('project.pst')
    assert p.pestmode == 'estimation' and p.nprior == 0
    assert p.counts['npar'] == 4 and p.counts['nobs'] == 4 and p.counts['npargp'] == 2
    assert p.par.set_index('PARNME').loc['hk3', 'TIETO'] == 'hk1'
    assert p.pestpp == [('overdue_resched_fac', '1.15'), ('max_run_fail', '3')]
    assert p.comments == ['created by makepst init']
    # the same workbook in regularisation mode picks up the PRIOR / WEIGHT columns
    args = shlex.split(cmd)[1:]
    args[1:3] = ['regul.pst', 'regul']
    main(args)
    assert read_pst('regul.pst').nprior == 2


def test_init_parrep_from_workbook(starter):
    book, cmd = starter
    par = book.parent / 'x.par'
    par.write_text('single point\nhk1 11 1 0\nhk2 22 1 0\nhk3 11 1 0\nrch 0.002 1 0\n')
    main(['parrep', str(book), str(par), 'rep.pst', '--set', 'noptmax=0'])
    p = read_pst('rep.pst')
    assert p.par.set_index('PARNME').loc['hk2', 'PARVAL1'] == 22 and p.control['noptmax'] == 0


def test_init_workbook_structure(starter):
    book, _ = starter
    wb = openpyxl.load_workbook(book)
    assert wb.sheetnames == ['CONTROL', 'PARGP', 'PAR', 'OBS', 'PRIOR', 'IO', 'PP', 'NOTES', 'BUILD']
    for ws in wb.worksheets:
        assert ws.freeze_panes == 'A2'
        assert ws.cell(1, 1).font.bold
    par = wb['PAR']
    hdr = [c.value for c in par[1]]
    assert hdr[:10] == 'PARNME PARTRANS PARCHGLIM PARVAL1 PARLBND PARUBND PARGP SCALE OFFSET DERCOM'.split()
    assert par.cell(1, hdr.index('PARTRANS') + 1).comment is not None
    lists = {dv.formula1: str(dv.sqref) for dv in par.data_validations.dataValidation}
    assert '"log,none,fixed,tied"' in lists and lists['"log,none,fixed,tied"'].startswith('B2')
    assert '"factor,relative"' in lists
    io_lists = [dv.formula1 for dv in wb['IO'].data_validations.dataValidation]
    assert '"cmd,tpl,ins"' in io_lists


def test_init_control_sheet_complete(starter):
    book, _ = starter
    ctl = pd.read_excel(book, 'CONTROL')
    assert list(ctl.columns) == ['LINE', 'NAME', 'DEFAULT', 'VALUE', 'DESCRIPTION']
    every = [k for s in ALL_SECTIONS for k in s.fields]
    assert list(ctl['NAME']) == every
    assert ctl['DESCRIPTION'].notna().all() and set(every) <= set(DESCRIPTIONS)
    assert ctl.set_index('NAME').loc['pestmode', 'VALUE'] == 'estimation'
    assert ctl.set_index('NAME').loc['noptmax', 'DEFAULT'] == 5
    p = Pst()
    p.set_control(ctl)                                   # only pestmode is set; nothing else, no warnings
    assert p.control == {} and p.pestmode == 'estimation'
    wb = openpyxl.load_workbook(book)
    ws = wb['CONTROL']
    flagged = {str(dv.sqref).split(':')[0] for dv in ws.data_validations.dataValidation}
    names = {ws.cell(r, 2).value: r for r in range(2, ws.max_row + 1)}
    assert f"D{names['lamforgive']}" in flagged and f"D{names['noptmax']}" not in flagged
    assert all(f"D{names[k]}" not in flagged for k in COMPUTED - {"pestmode"})   # pestmode has a drop-down


def test_init_refuses_to_overwrite(starter):
    book, _ = starter
    with pytest.raises(SystemExit, match='--force'):
        main(['init', str(book)])
    main(['init', str(book), '--force', '--name', 'other'])
    assert 'other.pst' in pd.read_excel(book, 'BUILD')['COMMAND'][0]


def test_empty_prior_sheet_keeps_mode():
    p = Pst('estimation')
    p.add_prior(pd.DataFrame(columns=['PINME', 'EQ', 'WEIGHT', 'OBGNME']))
    assert p.pestmode == 'estimation'
    p.add_prior(pd.DataFrame({'PINME': ['a'], 'EQ': ['1.0 * x = 1'], 'WEIGHT': [1], 'OBGNME': ['r']}))
    assert p.pestmode == 'regularisation'
