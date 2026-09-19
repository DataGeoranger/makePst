"""`makepst init`: a starter workbook that builds a valid control file as it is.

Every sheet has the headers `build` expects, a few example rows, header comments, drop-down
lists for the fields PEST restricts to a fixed vocabulary, and frozen header rows. The BUILD
sheet holds the command that builds it.
"""
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .sections import ALL_SECTIONS, COMPUTED

# one line per control variable; the PEST manual has the full story
DESCRIPTIONS = {
    'rstfle': 'restart / norestart: write restart files',
    'pestmode': 'estimation / regularisation / prediction / pareto (the command line overrides)',
    'npar': 'number of parameters - computed',
    'nobs': 'number of observations - computed',
    'npargp': 'number of parameter groups - computed',
    'nprior': 'number of prior information equations - computed',
    'nobsgp': 'number of observation groups - computed',
    'maxcompdim': 'max size of compressed Jacobian (0 = no compression)',
    'derzerolim': 'derivatives below this are stored as zero in a compressed Jacobian',
    'ntplfle': 'number of template files - computed',
    'ninsfle': 'number of instruction files - computed',
    'precis': 'single / double: precision of parameter values written to model input files',
    'dpoint': 'point / nopoint: keep the decimal point when writing values',
    'numcom': 'number of model command lines',
    'jacfile': '1 if the model writes derivatives itself (derivatives command line section)',
    'messfile': '1 to write a PEST-to-model message file',
    'obsreref': 'obsreref / noobsreref: observation re-referencing',
    'rlambda1': 'initial Marquardt lambda',
    'rlamfac': 'factor by which lambda is adjusted (or negative for PEST_HP lambda scheme)',
    'phiratsuf': 'phi ratio at which a lambda is accepted as sufficient',
    'phiredlam': 'relative phi reduction below which lambda testing stops',
    'numlam': 'max lambdas tested per iteration',
    'jacupdate': 'iterations of Broyden Jacobian update (0 = none, 999 = always)',
    'lamforgive': 'lamforgive / nolamforgive: tolerate model failures during lambda testing',
    'derforgive': 'derforgive / noderforgive: tolerate model failures during derivative runs',
    'win_mrun_hours': 'PEST_HP: model run time window in hours',
    'uptestmin': 'PEST_HP: minimum number of upgrade tests',
    'uptestlim': 'PEST_HP: maximum number of upgrade tests',
    'relparmax': 'max relative change of a relative-limited parameter per iteration',
    'facparmax': 'max factor change of a factor-limited parameter per iteration',
    'facorig': 'parameter below facorig * initial value is treated as facorig * initial value',
    'iboundstick': 'iterations a parameter sticks to a bound it hit',
    'upvecbend': '1 to bend the upgrade vector at bounds',
    'absparmax': 'absparmax(n)=value: max absolute change for absolute(n)-limited parameters',
    'phiredswh': 'relative phi reduction that switches from 2- to 3-point derivatives',
    'noptswitch': 'iteration before which no switch to higher-order derivatives',
    'splitswh': 'phi ratio that triggers split-slope derivatives',
    'doaui': 'aui / noaui: automatic user intervention',
    'dosenreuse': 'senreuse / nosenreuse: sensitivity reuse',
    'boundscale': 'boundscale / noboundscale: scale parameters by bound width',
    'noptmax': 'max optimisation iterations; 0 = one run, -1 / -2 = Jacobian only',
    'phiredstp': 'stop when phi improves by less than this over nphistp iterations',
    'nphistp': 'iterations for the phiredstp test',
    'nphinored': 'stop after this many iterations without phi reduction',
    'relparstp': 'stop when max relative parameter change is below this for nrelpar iterations',
    'nrelpar': 'iterations for the relparstp test',
    'phistopthresh': 'stop when phi falls below this',
    'lastrun': '1 to run the model once more with the final parameters',
    'phiabandon': 'abandon if phi exceeds this after the first iteration (-1 = never)',
    'icov': '1 to write the parameter covariance matrix',
    'icor': '1 to write the parameter correlation matrix',
    'ieig': '1 to write eigenvectors / eigenvalues',
    'ires': '1 to write the resolution matrix',
    'jcosave': 'jcosave / nojcosave: save the Jacobian as case.jco',
    'verboserec': 'verboserec / noverboserec: full run record',
    'jcosaveitn': 'jcosaveitn / nojcosaveitn: Jacobian per iteration',
    'reisaveitn': 'reisaveitn / noreisaveitn: residuals per iteration',
    'parsaveitn': 'parsaveitn / noparsaveitn: parameters per iteration',
    'parsaverun': 'parsaverun / noparsaverun: parameters per model run',
    'rrfsave': 'rrfsave / norrfsave: run-results file',
    'svdmode': '0 off, 1 SVD of JtQJ, 2 SVD of Q^1/2 J',
    'maxsing': 'singular values kept before truncation (<= adjustable parameters)',
    'eigthresh': 'eigenvalue ratio for truncation; 5e-7 or higher',
    'eigwrite': '1 to write case.svd',
    'lsqrmode': '1 to solve with LSQR instead of SVD',
    'lsqr_atol': 'LSQR tolerance a',
    'lsqr_btol': 'LSQR tolerance b',
    'lsqr_conlim': 'LSQR condition number limit',
    'lsqr_itnlim': 'LSQR max iterations (about 4 x adjustable parameters)',
    'lsqrwrite': '1 to write LSQR output',
    'maxaui': 'max parameters held per AUI iteration',
    'auistartopt': 'iteration at which AUI may start',
    'noauiphirat': 'AUI is not used if phi ratio is below this',
    'auirestitn': 'iterations between AUI restarts',
    'auisensrat': 'sensitivity ratio for holding a parameter',
    'auiholdmaxchg': '1 to hold the parameter with the largest change',
    'auinumfree': 'parameters freed per iteration',
    'auiphiratsuf': 'phi ratio at which AUI stops for the iteration',
    'auiphirataccept': 'phi ratio accepted after AUI',
    'nauinoaccept': 'AUI iterations without acceptance before giving up',
    'basepestfile': 'SVD-assist: base control file',
    'basejacfile': 'SVD-assist: base Jacobian file',
    'svda_mulbpa': '1 to multiply base parameters by super-parameter factors',
    'svda_scaladj': '1 to scale super parameters by sensitivity',
    'svda_extsuper': 'super-parameter extension option',
    'svda_supdercalc': '1 to compute super-parameter derivatives',
    'svda_par_excl': '1 to exclude parameters with PARCHGLIM absolute',
    'phimlim': 'target measurement objective function',
    'phimaccept': 'accepted measurement objective function (about 1.05 x phimlim)',
    'fracphim': 'fraction of current phi used as phimlim when phi is far above it',
    'memsave': 'memsave / nomemsave',
    'wfinit': 'initial regularisation weight factor',
    'wfmin': 'minimum weight factor',
    'wfmax': 'maximum weight factor',
    'linreg': 'linreg / nonlinreg: all prior information is linear',
    'regcontinue': 'regcontinue / noregcontinue: continue if phimlim cannot be reached',
    'wffac': 'weight factor adjustment factor',
    'wftol': 'weight factor convergence tolerance',
    'iregadj': 'inter-regularisation group weight adjustment (0-5)',
    'noptregadj': 'iterations between regularisation weight adjustments',
    'regweightrat': 'max ratio of regularisation group weights',
    'regsingthresh': 'singular value threshold for iregadj 4/5',
}

SHEETS = {
    'CONTROL': {
        'headers': ['LINE', 'NAME', 'DEFAULT', 'VALUE', 'DESCRIPTION'],
        'comments': {
            'NAME': 'PEST control variable (any section). Blank VALUE = built-in default.',
            'VALUE': 'Only this column is read. Counts (npar, nobs, ...) are computed and ignored here.',
        },
    },
    'PARGP': {
        'headers': 'PARGPNME INCTYP DERINC DERINCLB FORCEN DERINCMUL DERMTHD'.split(),
        'rows': [('hk', 'relative', 0.01, 0, 'switch', 2, 'parabolic'),
                 ('rch', 'absolute', 0.001, 0, 'switch', 2, 'parabolic')],
        'lists': {'INCTYP': 'relative,absolute,rel_to_max', 'FORCEN': 'switch,always_2,always_3,always_5,switch_5',
                  'DERMTHD': 'parabolic,best_fit,outside_pts,minvar'},
        'comments': {'PARGPNME': 'Group name used in PAR!PARGP. Groups no parameter uses are dropped.',
                     'DERINC': 'Derivative increment (relative or absolute per INCTYP).',
                     'FORCEN': 'switch: 2-point derivatives, then 3-point when phiredswh is reached.'},
    },
    'PAR': {
        'headers': 'PARNME PARTRANS PARCHGLIM PARVAL1 PARLBND PARUBND PARGP SCALE OFFSET DERCOM TIETO PRIOR WEIGHT'.split(),
        'rows': [('hk1', 'log', 'factor', 10, 0.1, 1000, 'hk', 1, 0, 1, None, 25, 1),
                 ('hk2', 'log', 'factor', 20, 0.1, 1000, 'hk', 1, 0, 1, None, 'hk1', 1),
                 ('hk3', 'tied', 'factor', 10, 0.1, 1000, 'hk', 1, 0, 1, 'hk1', None, None),
                 ('rch', 'none', 'relative', 0.001, 0.0001, 0.01, 'rch', 1, 0, 1, None, None, None)],
        'lists': {'PARTRANS': 'log,none,fixed,tied', 'PARCHGLIM': 'factor,relative'},
        'comments': {
            'PARNME': 'Unique, <= 12 characters for PEST (200 for PEST++); case is ignored.',
            'PARTRANS': 'log / none / fixed / tied. Tied needs TIETO; tied to a fixed parameter becomes fixed.',
            'PARCHGLIM': 'factor / relative, or absolute(n) with absparmax(n)= in CONTROL.',
            'PARVAL1': 'Initial value; must lie within PARLBND..PARUBND for adjustable parameters.',
            'TIETO': 'Parameter this one is tied to (PARTRANS = tied).',
            'PRIOR': 'Regularisation mode only: a preferred value, or another parameter name '
                     '(then p and that parameter are kept equal). Blank = no equation.',
            'WEIGHT': 'Weight of the PRIOR equation; <= 0 or blank = no equation.',
        },
    },
    'OBS': {
        'headers': 'OBSNME OBSVAL WEIGHT OBGNME'.split(),
        'rows': [('h_w1', 101.2, 1.0, 'head'), ('h_w2', 98.7, 1.0, 'head'),
                 ('h_w3', 95.0, 0.0, 'head'), ('q_gauge', 1250, 0.01, 'flow')],
        'comments': {'OBSNME': 'Unique, <= 20 characters for PEST; case is ignored.',
                     'WEIGHT': 'Formulas are fine here: build reads the computed value, update leaves formulas alone.',
                     'OBGNME': 'Observation group, <= 12 characters.'},
    },
    'PRIOR': {
        'headers': 'PINME EQ WEIGHT OBGNME'.split(),
        'rows': [],
        'comments': {'PINME': 'Explicit prior information. Usually not needed: PAR!PRIOR / WEIGHT generate equations.',
                     'EQ': 'e.g.  1.0 * log(hk1) - 1.0 * log(hk2) = 0'},
    },
    'IO': {
        'headers': ['TYPE', 'IN', 'OUT'],
        'rows': [('cmd', 'python model.py', None), ('tpl', 'model.tpl', 'model.in'),
                 ('ins', 'heads.ins', 'heads.out'), ('ins', 'flow.ins', 'flow.out')],
        'lists': {'TYPE': 'cmd,tpl,ins'},
        'comments': {'TYPE': 'cmd: model command line in IN. tpl: template IN writes model input OUT. '
                             'ins: instruction IN reads model output OUT.'},
    },
    'PP': {
        'headers': ['PP_VAR', 'VAL', 'NOTE'],
        'rows': [('overdue_resched_fac', 1.15, 'PEST++ options become ++name(value); blank VAL = not written'),
                 ('max_run_fail', 3, ''), ('lambdas', None, 'e.g. 0.1,1,10,100')],
        'comments': {'PP_VAR': 'PEST++ option name. Leave VAL blank to skip.'},
    },
    'NOTES': {
        'headers': ['COMMENT'],
        'rows': [('created by makepst init',)],
        'comments': {'COMMENT': 'Each cell becomes a "# ..." line at the top of the control file: keep a run history here.'},
    },
}

BUILD_LINES = [
    'makepst build {name}.pst estimation',
    '--set_ctl_xls {book},CONTROL',
    '--add_pargp_xls {book},PARGP',
    '--add_par_xls {book},PAR',
    '--add_obs_xls {book},OBS',
    '--add_prior_xls {book},PRIOR',
    '--add_io_xls {book},IO',
    '--add_pp_xls {book},PP',
    '--add_comment_xls {book},NOTES',
]

HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')
COMPUTED_FILL = PatternFill('solid', fgColor='EDEDED')


def _control_rows():
    for sec in ALL_SECTIONS:
        for k in sec.fields:
            yield (sec.name, k, sec.defaults.get(k), None, DESCRIPTIONS.get(k, ''))


def _style_headers(ws, headers, comments):
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        if h in comments:
            c.comment = Comment(comments[h], 'makepst')
            c.comment.width, c.comment.height = 320, 90
    ws.freeze_panes = 'A2'


def _widths(ws):
    for col in ws.columns:
        width = max(len(str(c.value)) for c in col if c.value is not None) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width, 8), 70)


def _list_validation(ws, col_letter, values, first_row, last_row, strict=True):
    dv = DataValidation(type='list', formula1=f'"{values}"', allow_blank=True,
                        showErrorMessage=strict, errorStyle='stop' if strict else 'warning',
                        errorTitle='makepst', error=f'expected one of: {values}')
    ws.add_data_validation(dv)
    dv.add(f'{col_letter}{first_row}:{col_letter}{last_row}')


def write_starter(path, name=None):
    """Write the starter workbook to `path`; returns the build command recorded in its BUILD sheet."""
    book = os.path.basename(path)
    name = name or os.path.splitext(book)[0]
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet('CONTROL')
    spec = SHEETS['CONTROL']
    _style_headers(ws, spec['headers'], spec['comments'])
    for r, row in enumerate(_control_rows(), 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        field = row[1]
        if field in COMPUTED:
            for c in range(1, 6):
                ws.cell(r, c).fill = COMPUTED_FILL
        flags = next((s.flags[field] for s in ALL_SECTIONS if field in s.flags), None)
        if flags:
            _list_validation(ws, 'D', ','.join(sorted(flags)), r, r)
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 2).value == 'pestmode':
            ws.cell(r, 4, 'estimation')
    ws.column_dimensions['E'].width = 80
    for c in ('A', 'B', 'C', 'D'):
        ws.column_dimensions[c].width = 16
    ws.column_dimensions['A'].width = 28

    for sheet, spec in SHEETS.items():
        if sheet == 'CONTROL':
            continue
        ws = wb.create_sheet(sheet)
        _style_headers(ws, spec['headers'], spec.get('comments', {}))
        for r, row in enumerate(spec.get('rows', []), 2):
            for c, v in enumerate(row, 1):
                ws.cell(r, c, v)
        for col, values in spec.get('lists', {}).items():
            letter = get_column_letter(spec['headers'].index(col) + 1)
            _list_validation(ws, letter, values, 2, 1000, strict=col != 'PARCHGLIM')
        _widths(ws)

    ws = wb.create_sheet('BUILD')
    _style_headers(ws, ['COMMAND'], {'COMMAND': 'The command that builds the control file from this workbook; '
                                                'makepst parrep <this workbook> ... uses it too.'})
    lines = [ln.format(name=name, book=book) for ln in BUILD_LINES]
    for r, ln in enumerate(lines, 2):
        ws.cell(r, 1, ln)
    ws.column_dimensions['A'].width = 60
    for w in wb.worksheets:
        w.sheet_view.zoomScale = 100
        for row in w.iter_rows(min_row=1, max_row=1):
            for c in row:
                c.alignment = Alignment(vertical='center')
    wb.save(path)
    return ' '.join(lines)
