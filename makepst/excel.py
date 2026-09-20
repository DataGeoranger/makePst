"""Workbook <-> Pst.

- load_table():      one sheet (or csv) as a DataFrame with upper-case headers
- to_workbook():     dump a Pst to a new workbook laid out so `build` can read it back
- update_workbook(): write values into an existing workbook, matched by PARNME / OBSNME,
                     leaving every other cell (formulas, macros, helper columns) alone
"""
import fnmatch
import os
import warnings

import pandas as pd

from .phi import ies_phi, phi_by_group
from .pst import OBS_COLS, PAR_COLS, PRIOR_COLS, Pst, read_obs_ensemble, read_par, read_res
from .sections import ALL_SECTIONS, COMPUTED, fmt

XL_ENGINE = 'openpyxl'


def load_table(spec):
    """'book.xlsx,SHEET' or 'table.csv' -> DataFrame; string headers upper-cased."""
    if ',' in spec:
        path, sheet = spec.rsplit(',', 1)
        engine = None if path.lower().endswith('.xls') else XL_ENGINE
        df = pd.read_excel(path, sheet, engine=engine)
    else:
        df = pd.read_csv(spec)
    df.columns = [c.strip().upper() if isinstance(c, str) else c for c in df.columns]
    return df


def expand_spec(spec):
    """'book.xlsx,PAR_*' -> ['book.xlsx,PAR_HK', 'book.xlsx,PAR_SY', ...] in workbook order.

    A spec without glob characters is returned as is (csv paths never expand).
    """
    if ',' not in spec:
        return [spec]
    path, pattern = spec.rsplit(',', 1)
    if not any(c in pattern for c in '*?['):
        return [spec]
    engine = None if path.lower().endswith('.xls') else XL_ENGINE
    names = pd.ExcelFile(path, engine=engine).sheet_names
    hits = [n for n in names if fnmatch.fnmatch(n.lower(), pattern.lower())]
    if not hits:
        raise ValueError(f'no sheet in {path} matches {pattern!r}; sheets: {names}')
    return [f'{path},{n}' for n in hits]


# ---------------------------------------------------------------------- dump
def control_table(pst: Pst):
    rows = []
    for sec in ALL_SECTIONS:
        for k in sec.fields:
            if k == 'pestmode':
                val = pst.pestmode
            elif k in COMPUTED:
                val = None
            else:
                val = pst.control.get(k)
            default = pst.counts.get(k, sec.defaults.get(k))
            rows.append((sec.name, k, default, val))
    return pd.DataFrame(rows, columns=['LINE', 'NAME', 'DEFAULT', 'VALUE'])


def _sheet_name(prefix, group):
    return f'{prefix}_{group}'[:31].upper()


def to_workbook(pst: Pst, path, split=False):
    """Write every table of `pst` to `path`; returns the `build` command that reads it back."""
    book = os.path.basename(path)
    par = pst.par[PAR_COLS + [c for c in ('TIETO',) if c in pst.par]]
    tables = {'CONTROL': control_table(pst), 'PARGP': pst.pargp}
    if split:
        for g, df in par.groupby('PARGP', sort=False):
            tables[_sheet_name('PAR', g)] = df
        for g, df in pst.obs.groupby('OBGNME', sort=False):
            tables[_sheet_name('OBS', g)] = df
    else:
        tables['PAR'] = par
        tables['OBS'] = pst.obs[OBS_COLS]
    if pst.nprior:
        tables['PRIOR'] = pst.prior[PRIOR_COLS]
    if pst.obs_cov:
        tables['OBSGP'] = pd.DataFrame({'OBGNME': pst.obsgp, 'COVFILE': [pst.obs_cov.get(g, '') for g in pst.obsgp]})
    tables['IO'] = pd.DataFrame([('cmd', c, '') for c in pst.cmd]
                                + [('tpl', a, b) for a, b in pst.tpl]
                                + [('ins', a, b) for a, b in pst.ins], columns=['TYPE', 'IN', 'OUT'])
    if pst.pestpp:
        tables['PP'] = pd.DataFrame(pst.pestpp, columns=['PP_VAR', 'VAL'])
    if pst.comments:
        tables['NOTES'] = pd.DataFrame({'COMMENT': pst.comments})

    cmd = [f'makepst build NEW.pst {pst.pestmode}', f'--set_ctl_xls "{book},CONTROL"']
    for name in tables:
        kind = name.split('_')[0].lower()
        if kind in ('pargp', 'par', 'obs', 'obsgp', 'prior', 'io', 'pp'):
            cmd.append(f'--add_{kind}_xls "{book},{name}"')
    if 'NOTES' in tables:
        cmd.append(f'--add_comment_xls "{book},NOTES"')
    tables['BUILD'] = pd.DataFrame({'COMMAND': cmd})

    with pd.ExcelWriter(path, engine=XL_ENGINE) as xw:
        for name, df in tables.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for i, c in enumerate(df.columns, 1):
                width = max([len(str(c))] + [len(fmt(v)) for v in df[c].head(500)]) + 2
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(width, 80)
    print(f'workbook written to {path}')
    return ' '.join(cmd)


# ---------------------------------------------------------------------- update
class _OpenpyxlBook:
    """No Excel needed; keeps VBA but drops charts/images and does not recalculate formulas."""

    def __init__(self, path):
        import openpyxl
        vba = path.lower().endswith('.xlsm')
        self.wb = openpyxl.load_workbook(path, keep_vba=vba)                    # formulas, written back
        self.cached = openpyxl.load_workbook(path, keep_vba=vba, data_only=True)  # last values Excel computed

    def sheets(self):
        return self.wb.sheetnames

    def read(self, sheet):
        return [list(r) for r in self.cached[sheet].iter_rows(values_only=True)]

    def index(self, sheet, key):
        """Return the header and (row, normalized key) pairs without materializing the sheet."""
        ws = self.cached[sheet]
        header = [str(c.value).strip().upper() if c.value is not None else '' for c in ws[1]]
        if key not in header:
            return header, []
        col = header.index(key) + 1
        names = [(row, fmt(ws.cell(row, col).value).lower()) for row in range(2, ws.max_row + 1)]
        return header, names

    def formulas(self, sheet):
        # a str starting with '=' is an ordinary formula; openpyxl gives ArrayFormula objects for {=...}
        return [[(isinstance(v, str) and v.startswith('=')) or (v is not None and not isinstance(v, (str, int, float, bool)))
                 for v in r] for r in self.wb[sheet].iter_rows(values_only=True)]

    def formula_cells(self, sheet, columns):
        ws = self.wb[sheet]
        cells = set()
        for col in columns:
            for row in range(2, ws.max_row + 1):
                value = ws.cell(row, col).value
                if ((isinstance(value, str) and value.startswith('='))
                        or (value is not None and not isinstance(value, (str, int, float, bool)))):
                    cells.add((row, col))
        return cells

    def write(self, sheet, row, col, value):
        """Write one cell; False when the cell belongs to an array formula and cannot be written."""
        cell = self.wb[sheet].cell(row, col)
        if cell.value is not None and not isinstance(cell.value, (str, int, float, bool)):
            return False
        cell.value = value
        return True

    def write_many(self, sheet, changes):
        refused = 0
        for row, col, value in changes:
            if not self.write(sheet, row, col, value):
                refused += 1
        return refused

    def write_table(self, sheet, df):
        if sheet in self.wb.sheetnames:
            self.wb.remove(self.wb[sheet])
        ws = self.wb.create_sheet(sheet)
        ws.append([str(c) for c in df.columns])
        for row in df.itertuples(index=False):
            ws.append([None if fmt(v) == '' else (v.item() if hasattr(v, 'item') else v) for v in row])

    def save(self, path):
        self.wb.save(path)

    def close(self):
        pass


class _XlwingsBook:
    """Drives Excel itself: everything in the workbook is preserved and recalculated."""

    def __init__(self, path):
        import xlwings as xw
        self.app = xw.App(visible=False, add_book=False)
        self.app.display_alerts = False          # never block on an Excel dialog; refused writes are counted
        self.app.screen_updating = False
        self._calculation = None
        self._manual_calculation = False
        try:
            self._enable_events = self.app.enable_events
        except (AttributeError, OSError):
            self._enable_events = True
        try:
            self.app.enable_events = False
            self.wb = self.app.books.open(os.path.abspath(path))
            try:
                self._calculation = self.app.calculation
                self.app.calculation = 'manual'
                self._manual_calculation = True
            except Exception:
                pass                            # batching is still fast when Excel controls calculation
        except Exception:
            self.app.quit()
            raise

    def sheets(self):
        return [s.name for s in self.wb.sheets]

    def read(self, sheet):
        return self._grid(self.wb.sheets[sheet].used_range.value)

    def index(self, sheet, key):
        """Read only the header and key column; COM cost then scales with rows, not used cells."""
        ws = self.wb.sheets[sheet]
        used = ws.used_range
        last_row, last_col = used.last_cell.row, used.last_cell.column
        raw_header = ws.range((1, 1), (1, last_col)).value
        header_values = raw_header if isinstance(raw_header, list) else [raw_header]
        if header_values and isinstance(header_values[0], list):
            header_values = header_values[0]
        header = [str(c).strip().upper() if c is not None else '' for c in header_values]
        if key not in header:
            return header, []
        col = header.index(key) + 1
        if last_row < 2:
            return header, []
        raw = ws.range((2, col), (last_row, col)).value
        values = raw if isinstance(raw, list) else [raw]
        if values and isinstance(values[0], list):
            values = [r[0] for r in values]
        return header, [(row, fmt(value).lower()) for row, value in enumerate(values, 2)]

    def formulas(self, sheet):
        # .formula shows only the anchor of a dynamic-array spill or a {=...} block as a formula;
        # the other cells are caught in write() through HasArray / HasSpill
        return [[isinstance(v, str) and v.startswith('=') for v in r]
                for r in self._grid(self.wb.sheets[sheet].used_range.formula)]

    def formula_cells(self, sheet, columns):
        ws = self.wb.sheets[sheet]
        last_row = ws.used_range.last_cell.row
        cells = set()
        if last_row < 2:
            return cells
        for col in columns:
            raw = ws.range((2, col), (last_row, col)).formula
            values = raw if isinstance(raw, list) else [raw]
            if values and isinstance(values[0], list):
                values = [r[0] for r in values]
            for row, value in enumerate(values, 2):
                if isinstance(value, str) and value.startswith('='):
                    cells.add((row, col))
        return cells

    @staticmethod
    def _grid(v):
        if v is None:
            return [[]]
        if isinstance(v, list):
            return v if v and isinstance(v[0], list) else [v]
        return [[v]]

    def write(self, sheet, row, col, value):
        """Write one cell; False when Excel would refuse (part of an array formula or spill range)."""
        rng = self.wb.sheets[sheet].range((row, col))
        try:
            api = rng.api
            if api.HasArray or getattr(api, 'HasSpill', False):
                return False
            rng.value = value
            return True
        except Exception:                       # COM error: protected sheet, merged cell, ...
            return False

    def write_many(self, sheet, changes):
        """Write contiguous runs in one COM assignment, falling back when Excel rejects a run."""
        from collections import defaultdict

        by_col = defaultdict(list)
        for row, col, value in changes:
            by_col[col].append((row, value))
        refused = 0
        ws = self.wb.sheets[sheet]
        for col, items in by_col.items():
            items.sort()
            runs = []
            run = []
            for item in items:
                if run and item[0] != run[-1][0] + 1:
                    runs.append(run)
                    run = []
                run.append(item)
            if run:
                runs.append(run)
            for run in runs:
                start, end = run[0][0], run[-1][0]
                values = [[value] for _, value in run]
                try:
                    ws.range((start, col), (end, col)).value = values
                except Exception:
                    refused += sum(not self.write(sheet, row, col, value) for row, value in run)
        return refused

    def write_table(self, sheet, df):
        names = [s.name for s in self.wb.sheets]
        ws = self.wb.sheets[sheet] if sheet in names else self.wb.sheets.add(sheet, after=self.wb.sheets[-1])
        ws.clear()
        rows = [[str(c) for c in df.columns]]
        rows += [[None if fmt(v) == '' else (v.item() if hasattr(v, 'item') else v) for v in r]
                 for r in df.itertuples(index=False)]
        ws.range('A1').value = rows

    def save(self, path):
        self.app.calculate()                     # one recalculation after all batched writes
        self.wb.save(os.path.abspath(path))

    def close(self):
        try:
            self.wb.close()
        finally:
            try:
                if self._manual_calculation and self._calculation is not None:
                    self.app.calculation = self._calculation
                self.app.enable_events = self._enable_events
            except Exception:
                pass
            self.app.quit()


def _open(path, backend):
    if backend is None:
        try:
            import xlwings  # noqa: F401
            backend = 'xlwings'
        except ImportError:
            backend = 'openpyxl'
    return {'xlwings': _XlwingsBook, 'openpyxl': _OpenpyxlBook}[backend](path), backend


def _update_sheet(book, sheet, key, values, create=(), overwrite_formulas=False):
    """Write `values` (DataFrame indexed by lower-case key) into `sheet`, matched on column `key`.

    Columns of `values` missing from the sheet are created only if listed in `create`.
    Cells holding a formula are left alone unless `overwrite_formulas`; cells Excel refuses
    (array formulas, spill ranges) are always left alone.
    Returns (rows updated, formula cells skipped, refused cells, names found in the sheet).
    """
    header, names = book.index(sheet, key)
    if not names:
        return 0, 0, 0, set()
    seen = {name for _, name in names if name}
    matched = [(i, name) for i, name in names if name in values.index]
    if not matched:
        return 0, 0, 0, seen
    cols = {}
    for c in values.columns:
        if c in header:
            cols[c] = header.index(c) + 1
        elif c in create:
            header.append(c)
            cols[c] = len(header)
            book.write(sheet, 1, cols[c], c)
    formulas = book.formula_cells(sheet, cols.values()) if not overwrite_formulas else set()
    hits = skipped = refused = 0
    changes = []
    for i, name in matched:
        hits += 1
        for c, col in cols.items():
            if (i, col) in formulas:
                skipped += 1
                continue
            v = values.at[name, c]
            changes.append((i, col, None if fmt(v) == '' else (v.item() if hasattr(v, 'item') else v)))
    refused += book.write_many(sheet, changes)
    return hits, skipped, refused, seen


def _sheet_selected(name, patterns):
    return not patterns or any(fnmatch.fnmatch(name.lower(), p.lower()) for p in patterns)


def update_workbook(path, par=None, obs=None, res=None, obs_csv=None, real=None, pst=None,
                    out=None, backend=None, par_cols=('PARVAL1',), obs_cols=('OBSVAL', 'WEIGHT'),
                    overwrite_formulas=False, sheets=None, groups=None, phi=True):
    """Update an existing workbook in place (or to `out`).

    par: a Pst, a DataFrame with PARNME, a .par file or an IES case.N.par.csv (realization `real`)
                                                              -> `par_cols` on every sheet with a PARNME column
    obs: a Pst or a DataFrame with OBSNME                     -> `obs_cols` on every sheet with an OBSNME column
    res: a .res/.rei file path                                 -> MODELLED / RESIDUAL columns on those sheets
    obs_csv: an IES case.N.obs.csv (realization `real`)        -> MODELLED, plus RESIDUAL when `pst` is given
    sheets: glob patterns (case-insensitive); only matching worksheets are touched
    groups: parameter / observation group names; only rows in those groups are written
            (a .par file carries no groups: pass a DataFrame or Pst with PARGP instead)
    Cells that hold formulas are skipped unless `overwrite_formulas`.
    phi: with residuals (res, or obs_csv + pst) also write a PHI sheet with the objective function
         by observation group, and PHI_IES with the realization phis of an ensemble when
         case.phi.actual.csv is beside it.
    """
    groups = {g.strip().lower() for g in groups} if groups else None

    def frame(src, key, group_col, cols):
        if src is None:
            return None
        if isinstance(src, Pst):
            src = src.par if key == 'PARNME' else src.obs
        elif isinstance(src, str):
            src = read_par(src, real).reset_index()
        df = src.copy()
        df.columns = [str(c).upper() for c in df.columns]
        df[key] = df[key].astype(str).str.strip().str.lower()
        if groups:
            if group_col not in df:
                raise ValueError(f'groups requested but {key} rows carry no {group_col} column')
            df = df[df[group_col].astype(str).str.strip().str.lower().isin(groups)]
        return df.set_index(key)[[c for c in cols if c in df]]

    par_df = frame(par, 'PARNME', 'PARGP', par_cols)
    obs_df = frame(obs, 'OBSNME', 'OBGNME', obs_cols)
    res_df = residuals = None
    if res is not None and obs_csv is not None:
        raise ValueError('give res or obs_csv, not both')
    if res is not None:
        residuals = read_res(res)
    elif obs_csv is not None:
        residuals = read_obs_ensemble(obs_csv, real, pst)
    if residuals is not None:
        res_df = frame(residuals.reset_index(), 'NAME', 'GROUP', ('MODELLED', 'RESIDUAL'))
    phi_tables = {}
    if phi and residuals is not None and {'GROUP', 'RESIDUAL', 'WEIGHT'} <= set(residuals.columns):
        phi_tables['PHI'] = phi_by_group(residuals)
        if obs_csv is not None:
            reals = ies_phi(obs_csv)
            if reals is not None:
                phi_tables['PHI_IES'] = reals
    if par_df is None and obs_df is None and res_df is None:
        raise ValueError('nothing to update: give par, obs, res or obs_csv')

    book, backend = _open(path, backend)
    if backend == 'openpyxl':
        warnings.warn('openpyxl backend: macros are kept, but charts/images are dropped and '
                      'formulas are not recalculated until Excel opens the file')
    touched = {}
    in_book = {'PARNME': set(), 'OBSNME': set()}       # names the selected sheets hold, per key
    try:
        for sheet in book.sheets():
            if not _sheet_selected(sheet, sheets):
                continue
            n = k = b = 0
            for key, df, create in (('PARNME', par_df, ()), ('OBSNME', obs_df, ()),
                                    ('OBSNME', res_df, ('MODELLED', 'RESIDUAL'))):
                if df is not None:
                    hits, skipped, refused, seen = _update_sheet(book, sheet, key, df, create, overwrite_formulas)
                    n, k, b = n + hits, k + skipped, b + refused
                    in_book[key] |= seen
            if n:
                print(f'{sheet}: {n} rows matched' + (f', {k} formula cells left alone' if k else '')
                      + (f', {b} cells refused by Excel (array formula / spill range)' if b else ''))
                touched[sheet] = {'rows': n, 'formula_cells_skipped': k, 'cells_refused': b}
        for name, table in phi_tables.items():
            book.write_table(name, table)
            print(f'{name}: {len(table)} rows written')
            touched[name] = {'rows': len(table), 'formula_cells_skipped': 0, 'cells_refused': 0}
        book.save(out or path)
    finally:
        book.close()

    # the two lists never match exactly in practice; say what was left out on either side
    unmatched = {}
    for key, df, what in (('PARNME', par_df, 'parameters'), ('OBSNME', obs_df, 'observations'),
                          ('OBSNME', res_df, 'observations')):
        if df is None or not in_book[key]:
            continue
        missing = sorted(set(df.index) - in_book[key])
        extra = sorted(in_book[key] - set(df.index))
        if missing:
            warnings.warn(f'{len(missing)} {what} in the results have no row in the workbook and were not written: '
                          f'{missing[:5]}{"..." if len(missing) > 5 else ""}')
        if extra:
            warnings.warn(f'{len(extra)} {what} in the workbook are not in the results and were left unchanged: '
                          f'{extra[:5]}{"..." if len(extra) > 5 else ""}')
        unmatched[what] = {'not_in_workbook': len(missing), 'not_in_results': len(extra)}
    print(f'workbook saved to {out or path} ({backend})')
    return {'sheets': touched, 'backend': backend, 'unmatched': unmatched}
